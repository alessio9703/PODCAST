"""Estrazione del testo dal documento di partenza (PDF, txt, md, docx)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".text"}


class DocumentError(RuntimeError):
    """Documento mancante, illeggibile o in un formato non supportato."""


class DocumentNotExtractable(DocumentError):
    """Il file c'e' ed e' del formato giusto, ma non se ne ricava testo.

    `details["motivo"]` e' uno fra:
      - "scansione": PDF fatto di sole immagini (serve OCR o un PDF con testo);
      - "protetto":  PDF protetto da password;
      - "vuoto":     DOCX/testo senza contenuto;
      - "illeggibile": file corrotto o non decodificabile.
    La UI mostra lo stesso blocco ("Questo PDF non contiene testo") per tutti,
    con il testo di `error` che dice cosa fare.
    """

    def __init__(self, message: str, *, motivo: str, file: Path) -> None:
        super().__init__(message)
        self.details = {"motivo": motivo, "file": file.name}


class UnsupportedDocument(DocumentError):
    """Estensione non gestita."""

    def __init__(self, message: str, *, extension: str, supported: list[str]) -> None:
        super().__init__(message)
        self.details = {"estensione": extension, "ammesse": supported}


@dataclass
class Document:
    path: Path
    text: str
    kind: str
    pages: int | None = None

    @property
    def char_count(self) -> int:
        return len(self.text)

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    def truncated(self, max_chars: int) -> tuple[str, bool]:
        if max_chars and self.char_count > max_chars:
            return self.text[:max_chars], True
        return self.text, False


SUPPORTED_EXTENSIONS = [".pdf", ".docx", *sorted(TEXT_EXTENSIONS)]


def locate_document(
    path: str | Path,
    search_dir: str | Path | None = None,
    *,
    confined: bool = False,
) -> Path:
    """Risolve il documento.

    Se `path` e' **solo un nome di file** (nessuna cartella davanti) viene
    cercato prima dentro `search_dir` — la cartella `documenti/` del progetto —
    e solo dopo come percorso relativo alla directory corrente.

    Con `confined=True` (i tool, per default: e' quello che usa l'interfaccia
    web) il documento DEVE stare dentro `search_dir`, una volta risolti i link
    simbolici; un nome nudo si cerca solo li', mai nella cartella corrente;
    file nascosti e file senza un'estensione gestita sono rifiutati. Senza,
    un nome come `.env` veniva cercato nella cartella corrente e letto come
    testo. run.py e il subagent, che lavorano con percorsi scelti in locale,
    usano la modalita' libera (--trusted-paths nei tool).
    """
    raw = Path(str(path).strip().strip('"').strip("'")).expanduser()
    bare_name = raw.parent == Path(".")

    if confined:
        return _locate_confined(raw, bare_name, search_dir)

    if bare_name and search_dir is not None:
        candidate = Path(search_dir).expanduser() / raw.name
        if candidate.is_file():
            return candidate.resolve()

    resolved = raw.resolve()
    if resolved.is_file():
        return resolved

    if bare_name and search_dir is not None:
        raise DocumentError(
            f"Documento non trovato: ne' '{raw.name}' dentro "
            f"{Path(search_dir)}, ne' come percorso {resolved}."
        )
    raise DocumentError(f"Documento non trovato: {resolved}")


def _locate_confined(raw: Path, bare_name: bool, search_dir: str | Path | None) -> Path:
    if search_dir is None:
        raise DocumentError("Cartella dei documenti non configurata (paths.documents_dir).")
    root = Path(search_dir).expanduser().resolve()
    candidate = (root / raw.name) if bare_name else raw
    resolved = candidate.resolve()
    try:
        resolved.relative_to(root)
    except ValueError:
        raise DocumentError(
            f"Il documento deve stare nella libreria Documenti ({root.name}/)."
        ) from None
    if resolved.name.startswith(".") or resolved.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocument(
            f"'{raw.name}' non e' un documento utilizzabile. Formati gestiti: PDF, DOCX, "
            f"{', '.join(sorted(TEXT_EXTENSIONS))}.",
            extension=resolved.suffix.lower(),
            supported=SUPPORTED_EXTENSIONS,
        )
    if not resolved.is_file():
        raise DocumentError(f"Documento non trovato nella libreria Documenti: '{raw.name}'.")
    return resolved


def extract_text(
    path: str | Path, search_dir: str | Path | None = None, *, confined: bool = False
) -> Document:
    doc_path = locate_document(path, search_dir, confined=confined)
    suffix = doc_path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(doc_path)
    if suffix == ".docx":
        return _read_docx(doc_path)
    if suffix in TEXT_EXTENSIONS or suffix == "":
        return _read_plain(doc_path)
    raise UnsupportedDocument(
        f"Estensione '{suffix}' non supportata. Formati gestiti: PDF, DOCX, "
        f"{', '.join(sorted(TEXT_EXTENSIONS))}.",
        extension=suffix,
        supported=SUPPORTED_EXTENSIONS,
    )


def _read_pdf(path: Path) -> Document:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentError(
            "Per leggere i PDF serve il pacchetto 'pypdf': pip install pypdf"
        ) from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentNotExtractable(
            f"Impossibile aprire il PDF {path.name}: {exc}", motivo="illeggibile", file=path
        ) from exc

    if getattr(reader, "is_encrypted", False):
        # Un PDF cifrato con password utente vuota si apre; con una password
        # vera pypdf NON solleva qui: decrypt() restituisce "non decifrato"
        # (0) e l'errore arriverebbe solo leggendo le pagine, con un tipo
        # interno di pypdf. Si controlla il risultato.
        try:
            decrypted = reader.decrypt("")
        except Exception:  # noqa: BLE001
            decrypted = 0
        if not decrypted:
            raise DocumentNotExtractable(
                f"Il PDF {path.name} e' protetto da password: rimuovi la protezione "
                "o fornisci una versione sbloccata.",
                motivo="protetto",
                file=path,
            )

    chunks: list[str] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            content = page.extract_text() or ""
        except Exception:
            content = ""
        if content.strip():
            chunks.append(f"[pagina {number}]\n{content.strip()}")

    text = "\n\n".join(chunks).strip()
    if not text:
        raise DocumentNotExtractable(
            f"Dal PDF {path.name} non e' stato estratto alcun testo: probabilmente "
            "e' un documento scansionato (immagini). Serve una versione con testo "
            "selezionabile oppure un passaggio OCR.",
            motivo="scansione",
            file=path,
        )
    return Document(path=path, text=text, kind="pdf", pages=len(reader.pages))


def _read_docx(path: Path) -> Document:
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise DocumentError(
            "Per leggere i .docx serve il pacchetto 'python-docx', che e' fra le "
            "dipendenze obbligatorie: .venv/bin/pip install -r requirements.txt "
            "(tools/doctor.py dice cosa manca)."
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocumentNotExtractable(
            f"Impossibile aprire il DOCX {path.name}: {exc}", motivo="illeggibile", file=path
        ) from exc
    text ="\n".join(p.text for p in document.paragraphs if p.text.strip()).strip()
    if not text:
        raise DocumentNotExtractable(
            f"Il documento {path.name} non contiene testo.", motivo="vuoto", file=path
        )
    return Document(path=path, text=text, kind="docx")


def _read_plain(path: Path) -> Document:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = path.read_text(encoding=encoding).strip()
        except UnicodeDecodeError:
            continue
        if not text:
            raise DocumentNotExtractable(f"Il file {path.name} e' vuoto.", motivo="vuoto", file=path)
        return Document(path=path, text=text, kind="text")
    raise DocumentNotExtractable(
        f"Impossibile decodificare {path.name} come file di testo.", motivo="illeggibile", file=path
    )


__all__ = [
    "Document",
    "DocumentError",
    "DocumentNotExtractable",
    "SUPPORTED_EXTENSIONS",
    "UnsupportedDocument",
    "extract_text",
    "locate_document",
]
